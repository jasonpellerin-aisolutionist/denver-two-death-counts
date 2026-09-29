"""Excel workbook a reader can open without Tableau or Python.

Reported figures stay on their own sheet. They are not added to the file totals.
Navy headers only.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from twocounts import EXPORTS, REPORTS

NAVY = "1E3A5F"
SLATE = "64748B"
INK = "0F172A"
WHITE = "FFFFFF"

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color=WHITE)
TITLE_FONT = Font(bold=True, size=16, color=INK)
NOTE_FONT = Font(italic=True, size=10, color=SLATE)


def _header(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_frame(ws, frame: pd.DataFrame, start: int = 1) -> None:
    labels = list(frame.columns)
    _header(ws, start, labels)
    for r, row in enumerate(frame.itertuples(index=False), start=start + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=None if pd.isna(value) else value)
            cell.font = Font(color=INK)
    _widths(ws, [28] * len(labels))


def main() -> None:
    yearly = pd.read_csv(EXPORTS / "yearly.csv")
    reported = pd.read_csv(EXPORTS / "reported.csv")
    claims = pd.read_csv(EXPORTS / "claims.csv")
    checks = pd.read_csv(EXPORTS / "quality_checks.csv")
    streets = pd.read_csv(EXPORTS / "streets.csv")

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme["A1"] = "Denver's two traffic-death counts"
    readme["A1"].font = TITLE_FONT
    notes = [
        "The open crash file and the city's published death counts are different publications.",
        "Do not add a reported figure to a file total. They are not the same series.",
        "2026 is year to date.",
        "Pedestrian crashes continue in the file after pedestrian deaths in that file go to zero.",
        "Street mentions split an intersection into two names. A mention is not a crash on that street.",
        "The gap is the finding. It is not a claim that either office falsified a number.",
    ]
    for i, note in enumerate(notes, start=3):
        readme.cell(row=i, column=1, value=note).font = NOTE_FONT
    _widths(readme, [120])

    file_sheet = wb.create_sheet("Open file")
    _write_frame(file_sheet, yearly)
    chart = BarChart()
    chart.title = "Deaths in the open file"
    chart.y_axis.title = "Deaths"
    data = Reference(file_sheet, min_col=4, min_row=1, max_row=1 + len(yearly))
    cats = Reference(file_sheet, min_col=1, min_row=2, max_row=1 + len(yearly))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    file_sheet.add_chart(chart, "A10")

    reported_sheet = wb.create_sheet("Reported figures")
    _write_frame(reported_sheet, reported)
    claims_sheet = wb.create_sheet("Claims")
    _write_frame(claims_sheet, claims)
    street_sheet = wb.create_sheet("Street mentions")
    _write_frame(street_sheet, streets)
    check_sheet = wb.create_sheet("Quality checks")
    _write_frame(check_sheet, checks)

    out = REPORTS / "two_death_counts.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
